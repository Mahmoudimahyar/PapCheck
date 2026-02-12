"""Main entry point for DOCX parsing: parse_manuscript()."""

import logging
from pathlib import Path

import docx
from docx.document import Document as DocxDocument

from refcheck.models.reference import (
    ManuscriptSection,
    ParsedManuscript,
    Reference,
)
from refcheck.stages.parse_docx.citation_detector import (
    detect_citation_style,
    find_citations_in_text,
)
from refcheck.stages.parse_docx.field_codes import detect_field_codes
from refcheck.stages.parse_docx.reference_extractor import (
    is_reference_heading,
    parse_single_reference,
    split_reference_section,
)

logger = logging.getLogger(__name__)

# Paragraph styles that indicate bibliography/reference entries
_BIBLIOGRAPHY_STYLES = {
    "endnote bibliography",
    "bibliography",
    "reference",
    "references",
    "works cited",
}


def parse_manuscript(docx_path: Path) -> ParsedManuscript:
    """Parse a DOCX manuscript into structured data."""
    warnings: list[str] = []
    try:
        doc = docx.Document(str(docx_path))
    except Exception as exc:
        logger.error("Failed to open DOCX: %s", exc)
        return ParsedManuscript(
            filename=docx_path.name,
            warnings=[f"Failed to open DOCX: {exc}"],
        )

    typed_doc: DocxDocument = doc  # type: ignore[assignment,unused-ignore]
    sections = _extract_sections(typed_doc, warnings)
    references = _extract_references(typed_doc, sections, warnings)
    body_text = " ".join(s.text for s in sections)
    citation_style = detect_citation_style(body_text)
    has_field_codes = detect_field_codes(docx_path)
    sections_with_citations = _add_citations_to_sections(sections)

    return ParsedManuscript(
        filename=docx_path.name,
        sections=sections_with_citations,
        references=references,
        citation_style=citation_style,
        has_field_codes=has_field_codes,
        warnings=warnings,
    )


def _extract_sections(
    doc: DocxDocument, warnings: list[str],
) -> list[ManuscriptSection]:
    """Extract text organized by headings from the document."""
    sections: list[ManuscriptSection] = []
    current_heading: str | None = None
    current_paragraphs: list[str] = []

    for para in doc.paragraphs:
        style_name = para.style.name if para.style else ""
        text = para.text.strip()
        if "Heading" in style_name and text:
            if current_paragraphs:
                sections.append(ManuscriptSection(
                    heading=current_heading, text="\n".join(current_paragraphs),
                ))
            current_heading = text
            current_paragraphs = []
        elif text:
            current_paragraphs.append(text)

    if current_paragraphs:
        sections.append(ManuscriptSection(
            heading=current_heading, text="\n".join(current_paragraphs),
        ))
    if not sections:
        warnings.append("No text content found in document")
    return sections


def _extract_references(
    doc: DocxDocument,
    sections: list[ManuscriptSection],
    warnings: list[str],
) -> list[Reference]:
    """Find and parse references using multiple strategies."""
    # Strategy 1: Bibliography-styled paragraphs (EndNote, Zotero, etc.)
    bib_refs = _extract_from_bibliography_style(doc)
    if bib_refs:
        logger.info("Extracted %d refs from bibliography style", len(bib_refs))
        return bib_refs

    # Strategy 2: Reference section by heading
    ref_section = _find_reference_section(sections)
    if ref_section is not None:
        raw_refs = split_reference_section(ref_section.text)
        if raw_refs:
            refs = _parse_raw_refs(raw_refs)
            logger.info("Extracted %d refs from section heading", len(refs))
            return refs
        warnings.append("Reference section found but no references parsed")
        return []

    warnings.append("No reference section found")
    return []


def _extract_from_bibliography_style(doc: DocxDocument) -> list[Reference]:
    """Extract references from paragraphs with bibliography styles."""
    bib_paragraphs: list[str] = []
    for para in doc.paragraphs:
        style_name = (para.style.name if para.style else "").lower()
        text = para.text.strip()
        if text and any(kw in style_name for kw in _BIBLIOGRAPHY_STYLES):
            bib_paragraphs.append(text)

    if len(bib_paragraphs) < 2:
        return []
    return _parse_raw_refs(bib_paragraphs)


def _find_reference_section(
    sections: list[ManuscriptSection],
) -> ManuscriptSection | None:
    """Find the section containing references by heading."""
    for section in sections:
        if section.heading and is_reference_heading(section.heading):
            return section
    return None


def _parse_raw_refs(raw_texts: list[str]) -> list[Reference]:
    """Parse a list of raw reference strings into Reference models."""
    references: list[Reference] = []
    for idx, raw_text in enumerate(raw_texts, start=1):
        ref = parse_single_reference(idx, raw_text)
        references.append(ref)
    return references


def _add_citations_to_sections(
    sections: list[ManuscriptSection],
) -> list[ManuscriptSection]:
    """Find in-text citations in each section."""
    result: list[ManuscriptSection] = []
    for section in sections:
        citations = find_citations_in_text(section.text)
        result.append(ManuscriptSection(
            heading=section.heading, text=section.text, citations=citations,
        ))
    return result
