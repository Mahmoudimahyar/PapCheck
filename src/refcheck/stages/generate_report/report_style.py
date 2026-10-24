"""Colours and small paragraph helpers shared by the report sections."""

from typing import TYPE_CHECKING

from docx.shared import Pt, RGBColor

if TYPE_CHECKING:
    from docx.document import Document as DocxDocument

GREEN = RGBColor(0x16, 0xA3, 0x4A)
RED = RGBColor(0xDC, 0x26, 0x26)
AMBER = RGBColor(0xD9, 0x77, 0x06)
GRAY = RGBColor(0x9C, 0xA3, 0xAF)
VIOLET = RGBColor(0x8B, 0x5C, 0xF6)

VERDICT_COLORS: dict[str, RGBColor] = {
    "supported": GREEN,
    "partially_supported": AMBER,
    "not_supported": RED,
    "contradicted": RED,
    "cannot_verify": GRAY,
    "not_relevant": VIOLET,
}


def verdict_label(verdict: str) -> str:
    """"partially_supported" -> "PARTIALLY SUPPORTED"."""
    return verdict.upper().replace("_", " ")


def add_colored_line(
    doc: "DocxDocument", text: str, color: RGBColor, bold: bool = False,
    style: str | None = None,
) -> None:
    """A paragraph with coloured (optionally bold) text."""
    para = doc.add_paragraph(style=style)
    run = para.add_run(text)
    run.bold = bold
    run.font.color.rgb = color
    run.font.size = Pt(11)


def add_labeled(doc: "DocxDocument", label: str, text: str, style: str | None = None) -> None:
    """"Label: text" with the label in bold."""
    para = doc.add_paragraph(style=style)
    para.add_run(f"{label}: ").bold = True
    para.add_run(text)
