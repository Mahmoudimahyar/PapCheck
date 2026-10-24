"""Report section: problems with the reference list and the attached PDFs."""

from typing import TYPE_CHECKING

from refcheck.stages.generate_report.source_issues import ISSUE_LABELS, SourceIssue

if TYPE_CHECKING:
    from docx.document import Document as DocxDocument

SECTION_TITLE = "Bibliography & Source Issues"


def _line(issue: SourceIssue) -> str:
    if not issue.claim_ids:
        return issue.message
    return f"{issue.message} Affected claims: {', '.join(map(str, issue.claim_ids))}."


def add_source_issues_section(doc: "DocxDocument", issues: list[SourceIssue]) -> None:
    """Group the issues by kind; retractions keep their own section above."""
    shown = [i for i in issues if i.kind != "retracted"]
    if not shown:
        return
    doc.add_heading(SECTION_TITLE, level=2)
    doc.add_paragraph(
        f"{len(shown)} problem(s) with the reference list or the attached PDFs. Fix these "
        "first: the affected claims were checked against another paper, only the abstract, "
        "or a guessed reference, so their verdicts say little about the manuscript."
    )
    if len(shown) < len(issues):
        doc.add_paragraph("Retracted or corrected references are listed in their own "
                          "section above.")
    for kind, label in ISSUE_LABELS.items():
        group = [i for i in shown if i.kind == kind]
        if not group:
            continue
        doc.add_heading(f"{label} ({len(group)})", level=3)
        for issue in group:
            doc.add_paragraph(_line(issue), style="List Bullet")
