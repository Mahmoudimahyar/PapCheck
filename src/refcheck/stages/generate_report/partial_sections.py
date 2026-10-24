"""V5.4 report sections: partial findings by sub-type, and supported results with a note.

Partial verdicts used to be one flat "Minor Issues" list, although some of them
(a wrong number, a conclusion the source does not draw) are major. They are
now filed under one sub-section per type, in the taxonomy's report order, each
finding with the author action. Partials listed under "Check First" (disputed
or with figures missing from the source) stay there and are only referenced.
"""

from typing import TYPE_CHECKING

from refcheck.models.claim import Claim
from refcheck.models.partial_types import GUARD_LABELS
from refcheck.models.verification import VerificationResult
from refcheck.stages.generate_report.claim_signals import claim_signals
from refcheck.stages.generate_report.finding_detail import add_finding_detail
from refcheck.stages.generate_report.partial_detail import add_partial_legend
from refcheck.stages.generate_report.partial_findings import (
    guarded_results,
    guards_of,
    partial_counts,
    partial_key,
    severity_of,
    type_info,
)
from refcheck.stages.generate_report.report_style import GRAY, add_colored_line
from refcheck.stages.generate_report.result_signals import check_first

if TYPE_CHECKING:
    from docx.document import Document as DocxDocument

PARTIAL_HEADING = "Partially Supported Findings by Type"
GUARDED_HEADING = "Supported with a Note"
_MAX_LISTED = 8  # guarded results listed per guard
_SNIPPET = 120


def _ref(v: VerificationResult) -> str:
    return f"Claim {v.claim_id} / Reference {v.reference_id}"


def _severity_summary(members: list[VerificationResult], default: str) -> str:
    major = sum(1 for v in members if severity_of(v) == "major")
    if major in (0, len(members)):
        return "major" if major else default
    return f"{major} major, {len(members) - major} minor"


def _add_overview(doc: "DocxDocument", groups: dict[str, list[VerificationResult]]) -> None:
    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    for cell, header in zip(table.rows[0].cells, ("Type", "Severity", "Findings", "Meaning"),
                            strict=True):
        cell.text = header
    for key, members in groups.items():
        label, meaning, severity = type_info(key)
        row = table.add_row().cells
        row[0].text, row[2].text, row[3].text = label, str(len(members)), meaning
        row[1].text = _severity_summary(members, severity)


def _add_type_section(doc: "DocxDocument", key: str, members: list[VerificationResult],
                      claims_by_id: dict[int, Claim]) -> None:
    label, meaning, _ = type_info(key)
    doc.add_heading(f"{label} ({len(members)})", level=3)
    add_colored_line(doc, meaning, GRAY)
    flagged = [v for v in members if check_first(v)]
    if flagged:
        doc.add_paragraph("Also of this type, listed under Check First above: "
                          + ", ".join(_ref(v) for v in flagged))
    ordered = sorted((v for v in members if not check_first(v)),
                     key=lambda v: (severity_of(v) != "major", v.claim_id))
    for v in ordered:
        add_finding_detail(doc, v, claims_by_id)


def add_partial_findings(
    doc: "DocxDocument", verifications: list[VerificationResult],
    claims_by_id: dict[int, Claim],
) -> None:
    """Legend, per-type overview and one sub-section per partial type."""
    groups = {key: [v for v in verifications if partial_key(v) == key]
              for key, count in partial_counts(verifications).items() if count}
    if not groups:
        return
    doc.add_heading(PARTIAL_HEADING, level=2)
    add_partial_legend(doc)
    _add_overview(doc, groups)
    for key, members in groups.items():
        _add_type_section(doc, key, members, claims_by_id)


def _snippet(v: VerificationResult, claims_by_id: dict[int, Claim]) -> str:
    claim = claims_by_id.get(v.claim_id)
    text = (claim_signals(claim).cited_for or claim.extracted_claim) if claim else ""
    text = text if len(text) <= _SNIPPET else f"{text[:_SNIPPET].rstrip()}…"
    return f"{_ref(v)}: \"{text}\"" if text else _ref(v)


def add_supported_with_note(
    doc: "DocxDocument", verifications: list[VerificationResult],
    claims_by_id: dict[int, Claim],
) -> None:
    """Supported results whose concern a benign guard cleared, grouped by guard."""
    guarded = guarded_results(verifications)
    if not guarded:
        return
    doc.add_heading(GUARDED_HEADING, level=2)
    doc.add_paragraph(
        f"{len(guarded)} supported citation(s) raised a concern that was checked and cleared. "
        "No change is required; the notes may still help (e.g. citation placement).")
    for guard, (label, why) in GUARD_LABELS.items():
        members = [v for v in guarded if guard in guards_of(v)]
        if not members:
            continue
        doc.add_heading(f"{label} ({len(members)})", level=3)
        add_colored_line(doc, why, GRAY)
        for v in members[:_MAX_LISTED]:
            doc.add_paragraph(_snippet(v, claims_by_id), style="List Bullet")
        if len(members) > _MAX_LISTED:
            doc.add_paragraph(f"… and {len(members) - _MAX_LISTED} more.")
