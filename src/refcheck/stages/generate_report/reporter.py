"""Generate verification report as DOCX using python-docx."""

import logging
from pathlib import Path
from typing import TYPE_CHECKING

import docx
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.shared import Pt, RGBColor

from refcheck.models.pipeline import PipelineState
from refcheck.models.reference import Reference
from refcheck.stages.generate_report.cost_section import add_cost_section
from refcheck.stages.generate_report.report_sections import (
    add_critical_findings,
    add_executive_summary,
    add_methodology_note,
    add_minor_issues,
    add_retraction_section,
    add_verified_section,
)

if TYPE_CHECKING:
    from docx.document import Document as DocxDocument
    from docx.table import _Cell

logger = logging.getLogger(__name__)

_GREEN = RGBColor(0x16, 0xA3, 0x4A)
_RED = RGBColor(0xDC, 0x26, 0x26)
_AMBER = RGBColor(0xD9, 0x77, 0x06)
_GRAY = RGBColor(0x9C, 0xA3, 0xAF)


def generate_report(state: PipelineState, output_path: Path) -> Path:
    """Generate a DOCX report from pipeline state (MVP + V1 sections)."""
    doc = docx.Document()
    has_verifications = bool(state.verification_results)

    _add_title(doc, state, has_verifications)
    _add_summary(doc, state)
    add_retraction_section(doc, state.references)

    # V1: Add verification sections if available
    if has_verifications:
        claims_by_id = {c.id: c for c in state.claims}
        add_executive_summary(doc, state.verification_results)
        add_critical_findings(doc, state.verification_results, claims_by_id)
        add_minor_issues(doc, state.verification_results, claims_by_id)
        add_verified_section(doc, state.verification_results)

    _add_reference_table(doc, state.references, has_verifications)
    _add_not_found_section(doc, state.references)

    # V4: Cost and tier statistics
    if has_verifications:
        add_cost_section(doc, state.verification_results)

    # V1: Methodology note
    if has_verifications:
        add_methodology_note(doc, len(state.verification_results))

    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output_path))
    logger.info("Report saved to %s", output_path)
    return output_path


def _add_title(
    doc: "DocxDocument", state: PipelineState, v1: bool,
) -> None:
    """Add report title and metadata."""
    title = "RefCheck AI — Verification Report" if v1 else (
        "RefCheck AI — Reference Existence Report"
    )
    doc.add_heading(title, level=1)
    filename = state.manuscript.filename if state.manuscript else "Unknown"
    doc.add_paragraph(f"Manuscript: {filename}")
    doc.add_paragraph(f"Session: {state.session_id}")


def _add_summary(doc: "DocxDocument", state: PipelineState) -> None:
    """Add summary statistics section."""
    doc.add_heading("Summary", level=2)
    refs = state.references
    total = len(refs)
    found = sum(1 for r in refs if r.source_status == "found")
    not_found = sum(1 for r in refs if r.source_status == "not_found")
    api_error = sum(1 for r in refs if r.source_status == "api_error")
    with_pdf = sum(1 for r in refs if r.pdf_path is not None)

    doc.add_paragraph(f"Total references: {total}")
    doc.add_paragraph(f"Found in databases: {found}")
    doc.add_paragraph(f"Not found: {not_found}")
    doc.add_paragraph(f"API errors: {api_error}")
    doc.add_paragraph(f"PDFs available: {with_pdf}")


def _add_reference_table(
    doc: "DocxDocument",
    references: list[Reference],
    include_verdict: bool = False,
) -> None:
    """Add reference status table."""
    doc.add_heading("Reference Status", level=2)
    if not references:
        doc.add_paragraph("No references to display.")
        return

    cols = 5
    headers = ["#", "Title", "Authors", "Status", "PDF"]
    table = doc.add_table(rows=1, cols=cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"

    for idx, header in enumerate(headers):
        cell = table.rows[0].cells[idx]
        cell.text = header
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.bold = True
                run.font.size = Pt(9)

    for ref in references:
        row = table.add_row()
        row.cells[0].text = str(ref.id)
        row.cells[1].text = ref.title or ""
        row.cells[2].text = ", ".join(ref.authors[:2])
        row.cells[3].text = ref.source_status
        row.cells[4].text = _pdf_status_text(ref)
        _color_status_cell(row.cells[3], ref.source_status)


def _pdf_status_text(ref: Reference) -> str:
    """Human-readable PDF status for a reference."""
    if ref.pdf_path:
        return ref.pdf_source or "available"
    if ref.journal_url:
        return "paywalled"
    return "missing"


def _color_status_cell(cell: "_Cell", status: str) -> None:
    """Apply color to a status cell based on reference status."""
    color_map = {"found": _GREEN, "not_found": _RED, "api_error": _AMBER, "pending": _GRAY}
    color = color_map.get(status, _GRAY)
    for paragraph in cell.paragraphs:
        for run in paragraph.runs:
            run.font.color.rgb = color


def _add_not_found_section(
    doc: "DocxDocument",
    references: list[Reference],
) -> None:
    """Add section for not-found references with search details."""
    not_found = [r for r in references if r.source_status == "not_found"]
    if not not_found:
        return

    doc.add_heading("Not Found References", level=2)
    doc.add_paragraph(
        f"{len(not_found)} references could not be found in any database."
    )

    for ref in not_found:
        display = ref.title or ref.raw_text
        doc.add_paragraph(
            f"[{ref.id}] {display}",
            style="List Bullet",
        )
        if ref.doi:
            doc.add_paragraph(f"  Searched DOI: {ref.doi}")
        if ref.pmid:
            doc.add_paragraph(f"  Searched PMID: {ref.pmid}")
