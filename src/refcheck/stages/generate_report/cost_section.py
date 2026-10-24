"""Report section: how the V5 cascade decided each pair, and what it cost."""

import logging
from collections import Counter
from typing import TYPE_CHECKING

from refcheck.models.verification import VerificationResult

if TYPE_CHECKING:
    from docx.document import Document as DocxDocument

logger = logging.getLogger(__name__)

_DECIDED_BY_LABELS = {
    "screen": "Fast screen (auto-cleared, verifier co-signed)",
    "verifier": "Verifier",
    "judge": "Judge (opt-in escalation)",
    "none": "Not decided by a model (no source or error)",
}
_LAYER_LABELS = {
    "screen": "Screen (Decisions API)", "verifier": "Verifier",
    "verifier_expanded": "Verifier, expanded evidence",
    "verifier_recheck": "Verifier, evidence re-check (elements searched in the full text)",
    "judge": "Judge",
    "confirm": "Cross-check (blind, other model family)",
    "adjudicate": "Element check (each unsupported element of a partial verdict re-read)",
    "adjudicate_cocited": "Element check, co-cited references (remaining elements searched "
                          "in the full text of the other references cited in the sentence)",
    "adjudicate_slots": "Element check, claim slots (a time frame / mechanism / population "
                        "/ recommendation stated nowhere in the source re-read)",
    "slot_probe": "Slot probe (blind question about the cited span; experimental)",
}


def _pct(count: int, total: int) -> str:
    return f"{count} ({count / total * 100:.1f}%)" if total else str(count)


def add_cost_section(
    doc: "DocxDocument",
    verifications: list[VerificationResult],
) -> None:
    """Add cascade routing statistics (decided by, layers reached, cost) to the report."""
    if not verifications:
        return
    total = len(verifications)
    doc.add_heading("Verification Statistics (screen → verify → cross-check)", level=2)
    doc.add_paragraph(f"Claim-reference pairs: {total}")
    cost = sum(v.cost_usd for v in verifications)
    doc.add_paragraph(f"LLM cost for these pairs: ${cost:.2f}")
    review = sum(1 for v in verifications if v.needs_user_review)
    doc.add_paragraph(f"Flagged for human review: {_pct(review, total)}")
    errors = sum(1 for v in verifications if v.cannot_verify_reason == "llm_error")
    if errors:
        doc.add_paragraph(f"Could not be checked because of LLM/API errors: {_pct(errors, total)}")

    decided: Counter[str] = Counter(str(v.decided_by) for v in verifications)
    doc.add_heading("Decided by", level=3)
    for key, label in _DECIDED_BY_LABELS.items():
        if decided.get(key):
            doc.add_paragraph(f"{label}: {_pct(decided[key], total)}", style="List Bullet")

    reached = Counter(layer for v in verifications for layer in set(v.layer_path))
    if reached:
        doc.add_heading("Layers reached", level=3)
        for key, label in _LAYER_LABELS.items():
            if reached.get(key):
                doc.add_paragraph(f"{label}: {_pct(reached[key], total)}", style="List Bullet")
