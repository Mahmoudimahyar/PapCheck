"""One flagged claim–reference pair in the report.

Besides the verdict, the reader needs what the reference is cited for (in a
multi-citation sentence it is one clause, not the sentence), both views when
the verifier and the cross-check disagreed, and the warnings the pipeline
raised (fragment, ambiguous mapping, unusable PDF, secondary citation, numbers
missing from the source). Before V5.2 the report printed only the whole
sentence and the verifier's verdict, so a disputed severity call read as a
firm "not supported".
"""

from typing import TYPE_CHECKING

from refcheck.models.claim import Claim
from refcheck.models.verification import VerificationResult
from refcheck.stages.generate_report.claim_signals import FRAGMENT_TEXT, claim_signals
from refcheck.stages.generate_report.partial_detail import add_partial_lines
from refcheck.stages.generate_report.report_style import (
    AMBER,
    GRAY,
    VERDICT_COLORS,
    add_colored_line,
    add_labeled,
    verdict_label,
)
from refcheck.stages.generate_report.result_signals import (
    DisputeViews,
    LayerView,
    ResultSignals,
    result_signals,
)

if TYPE_CHECKING:
    from docx.document import Document as DocxDocument

SECONDARY_TEXT = ("Secondary citation: the supporting passage itself cites other work — "
                  "consider citing the primary study.")


def _warnings(claim: Claim | None, signals: ResultSignals) -> list[str]:
    cs = claim_signals(claim)
    warnings = [FRAGMENT_TEXT] if cs.fragment else []
    if cs.mapping_text:
        warnings.append(cs.mapping_text)
    if signals.source_problem_text:
        warnings.append(f"{signals.source_problem_text}; it was not used as the article.")
    return warnings


def _notes(signals: ResultSignals) -> list[str]:
    notes = [SECONDARY_TEXT] if signals.secondary_citation else []
    if signals.numbers_not_found:
        notes.append("Numbers in the claim found nowhere in the source: "
                     + ", ".join(signals.numbers_not_found))
    return notes


def _view_line(role: str, view: LayerView) -> str:
    model = f" ({view.model_name})" if view.model_name else ""
    text = f"{role}{model}: {verdict_label(view.verdict)}"
    return f"{text} — {view.explanation}" if view.explanation else text


def add_dispute(doc: "DocxDocument", dispute: DisputeViews) -> None:
    """Both models' verdicts and reasons, so a human can settle the dispute."""
    add_colored_line(doc, "Disputed — the two models disagreed; the verdict shown is "
                          "pending human review:", AMBER, bold=True)
    doc.add_paragraph(_view_line("Verifier", dispute.verifier), style="List Bullet")
    doc.add_paragraph(_view_line("Cross-check", dispute.cross_check), style="List Bullet")


def add_finding_detail(
    doc: "DocxDocument", result: VerificationResult, claims_by_id: dict[int, Claim],
) -> None:
    """Add the detail block of one finding."""
    claim = claims_by_id.get(result.claim_id)
    signals = result_signals(result)
    doc.add_paragraph(f"Claim {result.claim_id} / Reference {result.reference_id}")
    if claim:
        doc.add_paragraph(f"Manuscript: \"{claim.manuscript_text}\"")
        cited_for = claim_signals(claim).cited_for
        if cited_for:
            add_labeled(doc, "Cited for", f"\"{cited_for}\" (the part of the sentence this "
                                          "reference supports; other parts cite other work)")
        else:
            doc.add_paragraph(f"Claim: {claim.extracted_claim}")
    for warning in _warnings(claim, signals):
        add_colored_line(doc, f"Warning: {warning}", AMBER)

    para = doc.add_paragraph("Verdict: ")
    run = para.add_run(verdict_label(result.verdict))
    run.bold = True
    run.font.color.rgb = VERDICT_COLORS.get(result.verdict, GRAY)
    if signals.dispute is not None:
        add_dispute(doc, signals.dispute)
    add_partial_lines(doc, result)  # V5.4: type, missing elements, what to do, cleared
    for note in _notes(signals):
        add_labeled(doc, "Note", note)

    if result.evidence_quotes:
        doc.add_paragraph("Evidence from source:")
        for quote in result.evidence_quotes:
            doc.add_paragraph(f'  "{quote}"', style="List Bullet")
    if result.reasoning:
        doc.add_paragraph(f"Reasoning: {result.reasoning}")
    doc.add_paragraph(
        f"Confidence: {result.confidence:.0%} | "
        f"Tier: {result.tier} | "
        f"Coverage: {result.source_coverage}"
    )
    doc.add_paragraph("")  # Spacer
