"""Tests for duplicate reference detection (V2)."""

from pathlib import Path

import docx

from refcheck.models.reference import Reference
from refcheck.stages.parse_docx import parse_manuscript
from refcheck.stages.parse_docx.duplicate_detector import detect_duplicates


class TestDuplicateDetector:
    def test_same_doi_detected(self) -> None:
        """Same DOI -> duplicate detected."""
        refs = [
            Reference(id=1, title="Paper A", doi="10.1234/test"),
            Reference(id=2, title="Paper A variant", doi="10.1234/test"),
        ]
        updated, warnings = detect_duplicates(refs)
        assert updated[1].duplicate_of == 1
        assert len(warnings) == 1

    def test_similar_titles_detected(self) -> None:
        """Very similar titles -> duplicate detected."""
        refs = [
            Reference(
                id=1,
                title="Effect of Drug X on Mortality in Elderly Patients",
            ),
            Reference(
                id=2,
                title="Effects of Drug X on Mortality in Elderly Patients",
            ),
        ]
        updated, warnings = detect_duplicates(refs)
        assert updated[1].duplicate_of == 1

    def test_different_refs_no_false_duplicate(self) -> None:
        """Different references -> no false positive."""
        refs = [
            Reference(
                id=1,
                title="Drug X for Cancer Treatment",
                doi="10.1234/a",
            ),
            Reference(
                id=2,
                title="Gene Therapy for Heart Disease",
                doi="10.5678/b",
            ),
        ]
        updated, warnings = detect_duplicates(refs)
        assert all(r.duplicate_of is None for r in updated)
        assert len(warnings) == 0

    def test_same_author_year_similar_title(self) -> None:
        """Same first author + year + similar title -> duplicate."""
        refs = [
            Reference(
                id=1,
                title="Novel Approach to Drug Delivery",
                authors=["Smith J"],
                year=2020,
            ),
            Reference(
                id=2,
                title="A Novel Approach to Drug Delivery Systems",
                authors=["Smith J"],
                year=2020,
            ),
        ]
        updated, warnings = detect_duplicates(refs)
        assert updated[1].duplicate_of == 1


class TestDuplicateDetectorM5:
    """M5: "Wade 2024a" and "Wade 2024b" (same DOI + title) were both kept."""

    def test_same_doi_with_resolver_prefix_and_case(self) -> None:
        refs = [
            Reference(id=105, title="Synthetic graft engineering: current progress",
                      doi="10.3389/FIMMU.2024.000001", year=2024, year_suffix="a"),
            Reference(id=106, title="Synthetic graft engineering: current progress",
                      doi="https://doi.org/10.3389/fimmu.2024.000001.", year=2024,
                      year_suffix="b"),
        ]
        updated, warnings = detect_duplicates(refs)
        assert updated[1].duplicate_of == 105
        assert warnings == ["Reference [106] appears to be a duplicate of [105] (same DOI)"]

    def test_same_title_author_year_without_doi(self) -> None:
        refs = [
            Reference(id=1, title="FasL microgels induce acceptance of islet grafts",
                      authors=["Lee, J.", "Cole, M. M."], year=2022, year_suffix="a"),
            Reference(id=2, title="FasL Microgels Induce Acceptance of Islet Grafts.",
                      authors=["Lee, J.", "Cole, M. M."], year=2022, year_suffix="b"),
        ]
        assert detect_duplicates(refs)[0][1].duplicate_of == 1

    def test_n_d_copy_of_dated_reference(self) -> None:
        refs = [
            Reference(id=67, title="Periosteum cues to hasten allograft healing", year=2025,
                      authors=["March, A."], doi="10.1000/jbm.a.1"),
            Reference(id=68, title="Periosteum cues to hasten allograft healing",
                      authors=["March, A."], doi="10.1000/JBM.A.1"),
        ]
        assert detect_duplicates(refs)[0][1].duplicate_of == 67

    def test_report_years_in_title_are_not_duplicates(self) -> None:
        refs = [
            Reference(id=1, title="Registry 2017 Annual Data Report: Heart", year=2019),
            Reference(id=2, title="Registry 2020 Annual Data Report: Heart", year=2022),
        ]
        assert all(r.duplicate_of is None for r in detect_duplicates(refs)[0])

    def test_same_title_different_first_author_not_duplicate(self) -> None:
        refs = [
            Reference(id=1, title="Hydrogels for local immunomodulation in grafts",
                      authors=["Wang, Q."], year=2024),
            Reference(id=2, title="Hydrogels for local immunomodulation in grafts",
                      authors=["Zhou, P."], year=2024),
        ]
        assert all(r.duplicate_of is None for r in detect_duplicates(refs)[0])

    def test_book_chapters_sharing_a_book_doi_are_distinct(self) -> None:
        refs = [
            Reference(id=1, title="Allograft choice in knee ligament surgery", doi="10.1000/book"),
            Reference(id=2, title="Rehabilitation protocols after meniscus repair",
                      doi="10.1000/book"),
        ]
        assert all(r.duplicate_of is None for r in detect_duplicates(refs)[0])


class TestParseManuscriptFlagsDuplicates:
    def test_parse_manuscript_runs_duplicate_detection(self, tmp_path: Path) -> None:
        """detect_duplicates() existed but parse_manuscript() never called it."""
        doc = docx.Document()
        doc.add_paragraph("Grafts fail (Wade et al., 2024a; Wade et al., 2024b).")
        doc.add_heading("References", level=1)
        for letter in ("a", "b"):
            doc.add_paragraph(
                f"Wade, Q., Holt, Y., and Lund, L. (2024{letter}). Synthetic graft engineering: "
                "current progress and open problems. Front. Synth. 15. doi: 10.1000/f.2024.1",
            )
        path = tmp_path / "dup.docx"
        doc.save(str(path))
        parsed = parse_manuscript(path)
        assert [r.duplicate_of for r in parsed.references] == [None, 1]
        assert any("duplicate of [1]" in w for w in parsed.warnings)
