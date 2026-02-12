"""Unit tests for report generation."""

from pathlib import Path

import docx
import pytest

from refcheck.models.pipeline import PipelineState
from refcheck.models.reference import ParsedManuscript, Reference
from refcheck.stages.generate_report import generate_report


def _make_state(refs: list[Reference] | None = None) -> PipelineState:
    """Create a pipeline state with optional references."""
    manuscript = ParsedManuscript(filename="test.docx")
    return PipelineState(
        session_id="test_session",
        manuscript=manuscript,
        references=refs or [],
    )


class TestReporter:
    def test_rg01_generates_valid_docx(self, tmp_path: Path) -> None:
        """RG-01: Report generates a valid DOCX file."""
        refs = [
            Reference(id=1, title="Found Paper", source_status="found"),
            Reference(id=2, title="Not Found Paper", source_status="not_found"),
        ]
        state = _make_state(refs)
        output = tmp_path / "report.docx"

        result = generate_report(state, output)
        assert result.exists()
        # Verify it's a valid DOCX
        doc = docx.Document(str(result))
        assert len(doc.paragraphs) > 0

    def test_rg02_summary_counts_match(self, tmp_path: Path) -> None:
        """RG-02: Summary counts match input data."""
        refs = [
            Reference(id=1, title="Paper 1", source_status="found"),
            Reference(id=2, title="Paper 2", source_status="found"),
            Reference(id=3, title="Paper 3", source_status="not_found"),
            Reference(id=4, title="Paper 4", source_status="api_error"),
        ]
        state = _make_state(refs)
        output = tmp_path / "report.docx"

        generate_report(state, output)
        doc = docx.Document(str(output))
        full_text = "\n".join(p.text for p in doc.paragraphs)

        assert "Total references: 4" in full_text
        assert "Found in databases: 2" in full_text
        assert "Not found: 1" in full_text
        assert "API errors: 1" in full_text

    def test_rg04_empty_results(self, tmp_path: Path) -> None:
        """RG-04: Empty results produce a valid report."""
        state = _make_state([])
        output = tmp_path / "report.docx"

        result = generate_report(state, output)
        assert result.exists()
        doc = docx.Document(str(result))
        assert len(doc.paragraphs) > 0

    def test_rg05_large_report(self, tmp_path: Path) -> None:
        """RG-05: 150 references generates in reasonable time."""
        import time

        refs = [
            Reference(
                id=i,
                title=f"Paper {i}: A Study of Topic {i}",
                authors=[f"Author{i} A", f"Author{i} B"],
                source_status="found" if i % 3 != 0 else "not_found",
            )
            for i in range(1, 151)
        ]
        state = _make_state(refs)
        output = tmp_path / "report.docx"

        start = time.time()
        result = generate_report(state, output)
        elapsed = time.time() - start

        assert result.exists()
        assert elapsed < 60  # Must complete in under 60 seconds

    def test_not_found_section(self, tmp_path: Path) -> None:
        """Not-found section lists unfound references."""
        refs = [
            Reference(
                id=1, title="Missing Paper", source_status="not_found",
                doi="10.1234/missing",
            ),
        ]
        state = _make_state(refs)
        output = tmp_path / "report.docx"

        generate_report(state, output)
        doc = docx.Document(str(output))
        full_text = "\n".join(p.text for p in doc.paragraphs)
        assert "Missing Paper" in full_text
