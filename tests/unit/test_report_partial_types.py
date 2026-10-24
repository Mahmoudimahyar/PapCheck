"""V5.4 DOCX report: partial findings by sub-type with author actions, legend, and
supported results with a note grouped by guard. Synthetic data only."""

from pathlib import Path

import docx
import pytest

from refcheck.models.claim import Claim
from refcheck.models.partial_types import GUARD_LABELS, author_action
from refcheck.models.pipeline import PipelineState
from refcheck.models.reference import ParsedManuscript, Reference
from refcheck.models.verification import VerificationResult
from refcheck.stages.generate_report import generate_report
from refcheck.stages.generate_report.partial_detail import LEGEND_TITLE
from refcheck.stages.generate_report.partial_sections import GUARDED_HEADING, PARTIAL_HEADING

_NUMBERS = "Numbers not found in the source: 40 kPa — please check where they come from."


def _v(claim_id: int, verdict: str, **extra: object) -> VerificationResult:
    return VerificationResult.model_validate({
        "claim_id": claim_id, "reference_id": 1, "verdict": verdict, "confidence": 0.7,
        "reasoning": f"Synthetic reasoning {claim_id}.", **extra})


def _results() -> list[VerificationResult]:
    partial = "partially_supported"
    return [
        _v(1, partial, partial_type="certainty_overstated", partial_severity="minor"),
        _v(2, partial, partial_type="missing_element", unsupported_elements=["fast resorption"]),
        _v(3, partial, partial_type="specific_detail_mismatch", partial_severity="major",
           unsupported_elements=["12 weeks"]),
        _v(4, partial, partial_type="specific_detail_mismatch", partial_severity="major",
           reasoning=f"Stiffness differs. {_NUMBERS}"),  # check first
        _v(5, partial),  # legacy row without a type
        *[_v(10 + i, "supported", guards_applied=["cocited_share"]) for i in range(10)],
        _v(30, "supported", guards_applied=["overliteral_wording"]),
        _v(31, "supported"),
    ]


def _text(tmp_path: Path, results: list[VerificationResult]) -> tuple[str, docx.Document]:
    claims = [Claim(id=v.claim_id, extracted_claim=f"Synthetic claim {v.claim_id}",
                    reference_ids=[1]) for v in results]
    state = PipelineState(session_id="s", manuscript=ParsedManuscript(filename="m.docx"),
                          references=[Reference(id=1, title="Paper", source_status="found")],
                          claims=claims, verification_results=results)
    document = docx.Document(str(generate_report(state, tmp_path / "r.docx")))
    return "\n".join(p.text for p in document.paragraphs), document


@pytest.fixture
def report(tmp_path: Path) -> tuple[str, docx.Document]:
    return _text(tmp_path, _results())


def test_sections_follow_the_taxonomy_order(report: tuple[str, docx.Document]) -> None:
    text, _ = report
    order = ["Check First", PARTIAL_HEADING, LEGEND_TITLE, "Wrong or missing specifics (2)",
             "Missing pieces (1)", "Overstated certainty (1)", "Other partial findings (1)",
             "Verified References", GUARDED_HEADING]
    positions = [text.index(heading) for heading in order]
    assert positions == sorted(positions)


def test_findings_show_type_elements_and_author_action(report: tuple[str, docx.Document]) -> None:
    text, _ = report
    section = text[text.index("Wrong or missing specifics (2)"):text.index("Missing pieces (1)")]
    assert "Issue type: Wrong or missing specifics (major)" in section
    assert "Not stated in the source: 12 weeks" in section
    assert f"What to do: {author_action('specific_detail_mismatch', '12 weeks')}" in section
    assert "Also of this type, listed under Check First above: Claim 4 / Reference 1" in section
    assert "Synthetic claim 4" not in section  # listed once, under Check First
    missing = text[text.index("Missing pieces (1)"):text.index("Overstated certainty (1)")]
    assert "Issue type: Missing pieces (minor)" in missing  # default severity
    assert author_action("missing_element", "fast resorption") in missing


def test_overview_table_counts_types(report: tuple[str, docx.Document]) -> None:
    _, document = report
    rows = [[c.text for c in row.cells] for t in document.tables for row in t.rows]
    assert ["Type", "Severity", "Findings", "Meaning"] in rows
    specifics = next(r for r in rows if r[0] == "Wrong or missing specifics")
    assert specifics[1:3] == ["major", "2"]


def test_supported_with_a_note_groups_by_guard(report: tuple[str, docx.Document]) -> None:
    text, _ = report
    section = text[text.index(GUARDED_HEADING):]
    assert "11 supported citation(s)" in section
    assert f"{GUARD_LABELS['cocited_share'][0]} (10)" in section
    assert "… and 2 more." in section
    assert f"{GUARD_LABELS['overliteral_wording'][0]} (1)" in section
    assert 'Claim 30 / Reference 1: "Synthetic claim 30"' in section
    assert "Claim 31 " not in section


def test_sections_are_omitted_without_partials_or_guards(tmp_path: Path) -> None:
    text, _ = _text(tmp_path, [_v(1, "supported"), _v(2, "not_supported")])
    assert PARTIAL_HEADING not in text and GUARDED_HEADING not in text
    assert LEGEND_TITLE not in text
