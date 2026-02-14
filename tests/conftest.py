"""Shared test fixtures for RefCheck tests."""

import asyncio
import sys
from collections.abc import Generator
from pathlib import Path

import docx
import fitz
import pytest
from dotenv import load_dotenv
from sqlmodel import Session, SQLModel, create_engine

# Load .env file so API keys are available in tests
load_dotenv(Path(__file__).parent.parent / ".env")

# Fix "Event loop is closed" error on Windows with httpx/litellm async calls.
if sys.platform == "win32":
    asyncio.set_event_loop_policy(asyncio.WindowsSelectorEventLoopPolicy())

FIXTURES_DIR = Path(__file__).parent / "fixtures"
MANUSCRIPTS_DIR = FIXTURES_DIR / "manuscripts"
PDFS_DIR = FIXTURES_DIR / "pdfs"

# Import all DB models so SQLModel.metadata registers them
import refcheck.db.models  # noqa: F401, E402


@pytest.fixture
def test_engine():
    """Create an in-memory SQLite engine with all tables."""
    engine = create_engine(
        "sqlite:///:memory:",
        echo=False,
        connect_args={"check_same_thread": False},
    )
    SQLModel.metadata.create_all(engine)
    return engine


@pytest.fixture
def db_session(test_engine: object) -> Generator[Session, None, None]:
    """Provide a database session bound to a fresh in-memory DB."""
    with Session(test_engine) as session:  # type: ignore[arg-type]
        yield session


@pytest.fixture
def fixtures_dir() -> Path:
    """Path to the test fixtures directory."""
    return FIXTURES_DIR


@pytest.fixture
def manuscripts_dir() -> Path:
    """Path to the manuscript fixtures directory."""
    return MANUSCRIPTS_DIR


def create_test_docx(
    path: Path,
    ref_count: int = 10,
    include_dois: bool = True,
    include_pmids: bool = True,
    numbered_format: str = "bracket",
) -> Path:
    """Create a realistic test DOCX with numbered references."""
    doc = docx.Document()
    doc.add_heading("Test Manuscript: Nanomedicine", level=1)
    doc.add_heading("Abstract", level=2)
    doc.add_paragraph(
        "This study reviews recent advances in nanomedicine "
        "for treating osteoarthritis [1, 2]."
    )
    doc.add_heading("1. Introduction", level=2)
    doc.add_paragraph(
        "Osteoarthritis is a chronic condition [1-3]. "
        "Recent studies [4, 5] have shown that nanoparticles "
        "can deliver drugs directly to joints [6, 7]."
    )
    doc.add_heading("2. Methods", level=2)
    doc.add_paragraph(
        "We followed established protocols [8-10] for review."
    )
    doc.add_heading("References", level=2)
    _add_sample_refs(doc, ref_count, include_dois, include_pmids)
    doc.save(str(path))
    return path


def _add_sample_refs(
    doc: docx.Document, ref_count: int,
    include_dois: bool, include_pmids: bool,
) -> None:
    """Add sample references to a docx document."""
    samples = _get_sample_refs()
    for i in range(min(ref_count, len(samples))):
        authors, title, journal, year, doi, pmid = samples[i]
        ref_text = f"[{i + 1}] {authors}, {title}, {journal}, ({year})"
        if doi and include_dois:
            ref_text += f". doi:{doi}"
        if pmid and include_pmids:
            ref_text += f" PMID:{pmid}"
        doc.add_paragraph(ref_text + ".")

    for i in range(len(samples), ref_count):
        doc.add_paragraph(
            f"[{i + 1}] Author{i}, Study {i}, J Generic, "
            f"({2015 + i % 10})."
        )


def _get_sample_refs() -> (
    list[tuple[str, str, str, int, str | None, str | None]]
):
    """Return sample reference tuples."""
    return [
        ("A. Smith, B. Jones",
         "Drug X reduces inflammation in knee joints",
         "Lancet", 2020,
         "10.1016/S0140-6736(20)30001-1", "32456789"),
        ("C. Wilson, D. Lee",
         "CRISPR-based diagnostics",
         "Nature Methods", 2019,
         "10.1038/s41592-019-0001-1", "31234567"),
        ("E. Brown, F. Garcia",
         "Nanoparticle delivery systems",
         "Adv Drug Deliv Rev", 2021,
         "10.1016/j.addr.2021.01.001", None),
        ("G. Martinez",
         "Global burden of osteoarthritis",
         "Lancet Rheumatol", 2023, None, None),
        ("H. Taylor, I. Anderson",
         "Stem cell therapy for cartilage",
         "Arthritis Res Ther", 2018,
         "10.1186/s13075-018-0001-1", None),
        ("J. Thomas",
         "Exosome therapeutics for joints",
         "Biomaterials", 2022, None, "35678901"),
        ("K. White, L. Harris",
         "siRNA delivery to chondrocytes",
         "Mol Ther Nucleic Acids", 2021,
         "10.1016/j.omtn.2021.05.001", None),
        ("M. Clark",
         "Autophagy in joint degeneration",
         "Cell Death Dis", 2020, None, None),
        ("N. Lewis, O. Walker",
         "Curcumin nanoformulations",
         "Int J Pharm", 2019,
         "10.1016/j.ijpharm.2019.118001", "30987654"),
        ("P. Hall, Q. Young",
         "ML for drug delivery",
         "Adv Sci", 2024, None, None),
    ]


def create_test_pdf(
    path: Path, title: str,
    doi: str = "", abstract: str = "",
) -> Path:
    """Create a test PDF with title and optional DOI."""
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((72, 72), title, fontsize=18)
    text = abstract or "This paper presents novel findings."
    page.insert_text((72, 120), f"Abstract: {text}", fontsize=10)
    if doi:
        page.insert_text((72, 160), f"https://doi.org/{doi}", fontsize=9)
    doc.save(str(path))
    doc.close()
    return path
