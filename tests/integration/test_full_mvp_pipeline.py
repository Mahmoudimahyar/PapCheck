"""Integration test: Full MVP pipeline end-to-end."""

from pathlib import Path

import docx
import pytest

from refcheck.models.pipeline import PipelineState
from refcheck.models.reference import ParsedManuscript
from refcheck.stages.generate_report import generate_report
from refcheck.stages.match_pdfs import match_pdfs
from refcheck.stages.parse_docx import parse_manuscript
from tests.conftest import create_test_docx, create_test_pdf


class TestFullMVPPipeline:
    def test_full_pipeline(self, tmp_path: Path) -> None:
        """Full MVP pipeline: parse -> match -> report."""
        # Stage 1: Create and parse a test DOCX
        docx_path = create_test_docx(tmp_path / "test_paper.docx", ref_count=10)
        manuscript = parse_manuscript(docx_path)
        assert len(manuscript.references) == 10
        assert manuscript.citation_style == "numbered"

        # Stage 3: Create matching PDFs and match
        pdf_dir = tmp_path / "pdfs"
        pdf_dir.mkdir()

        # Create PDFs that match the first 5 references by title
        titles_and_dois = [
            (
                "Drug X reduces inflammation in knee joints",
                "10.1016/S0140-6736(20)30001-1",
            ),
            (
                "CRISPR-based diagnostics for early disease detection",
                "10.1038/s41592-019-0001-1",
            ),
            (
                "Nanoparticle delivery systems for targeted therapy",
                "10.1016/j.addr.2021.01.001",
            ),
            (
                "Global burden of osteoarthritis in aging populations",
                "",
            ),
            (
                "Mesenchymal stem cell therapy for cartilage repair",
                "10.1186/s13075-018-0001-1",
            ),
        ]

        for i, (title, doi) in enumerate(titles_and_dois):
            create_test_pdf(
                pdf_dir / f"paper_{i}.pdf",
                title=title,
                doi=doi,
            )

        results = match_pdfs(manuscript.references, pdf_dir)
        matched_ids = {
            r.reference_id for r in results if r.match_method != "unmatched"
        }
        assert len(matched_ids) >= 3  # At least some DOI + title matches

        # Stage 6: Generate report
        state = PipelineState(
            session_id="integration_test",
            manuscript=manuscript,
            references=manuscript.references,
            match_results=results,
        )
        report_path = tmp_path / "report.docx"
        generate_report(state, report_path)

        assert report_path.exists()
        report_doc = docx.Document(str(report_path))
        full_text = "\n".join(p.text for p in report_doc.paragraphs)
        assert "Total references: 10" in full_text
        assert "RefCheck AI" in full_text

    def test_parse_extracts_all_metadata(self, tmp_path: Path) -> None:
        """Parse stage extracts DOIs, PMIDs, and years correctly."""
        docx_path = create_test_docx(
            tmp_path / "meta.docx",
            ref_count=10,
            include_dois=True,
            include_pmids=True,
        )
        manuscript = parse_manuscript(docx_path)
        assert len(manuscript.references) == 10

        # Check DOIs extracted
        dois = [r for r in manuscript.references if r.doi]
        assert len(dois) >= 3

        # Check PMIDs extracted
        pmids = [r for r in manuscript.references if r.pmid]
        assert len(pmids) >= 2

        # Check years extracted
        years = [r for r in manuscript.references if r.year]
        assert len(years) >= 8

    def test_doi_matching_precision(self, tmp_path: Path) -> None:
        """DOI matching is precise - no false matches."""
        docx_path = create_test_docx(tmp_path / "test.docx", ref_count=5)
        manuscript = parse_manuscript(docx_path)

        pdf_dir = tmp_path / "pdfs"
        pdf_dir.mkdir()
        # Create a PDF with a DOI that matches ref 1
        create_test_pdf(
            pdf_dir / "match.pdf",
            title="Drug X Study",
            doi="10.1016/S0140-6736(20)30001-1",
        )
        # Create a PDF with a DOI that doesn't match any ref
        create_test_pdf(
            pdf_dir / "nomatch.pdf",
            title="Unrelated Paper",
            doi="10.9999/unrelated",
        )

        results = match_pdfs(manuscript.references, pdf_dir)
        doi_matches = [r for r in results if r.match_method == "doi"]
        assert len(doi_matches) == 1
        assert doi_matches[0].reference_id == 1

    def test_empty_pdf_directory(self, tmp_path: Path) -> None:
        """Match stage handles empty PDF directory gracefully."""
        docx_path = create_test_docx(tmp_path / "test.docx")
        manuscript = parse_manuscript(docx_path)
        pdf_dir = tmp_path / "empty_pdfs"
        pdf_dir.mkdir()

        results = match_pdfs(manuscript.references, pdf_dir)
        assert results == []

    def test_report_with_mixed_statuses(self, tmp_path: Path) -> None:
        """Report correctly shows mixed reference statuses."""
        from refcheck.models.reference import Reference

        refs = [
            Reference(
                id=1,
                title="Found Paper",
                source_status="found",
                pdf_path=Path("/tmp/test.pdf"),
                pdf_source="user_upload",
            ),
            Reference(id=2, title="Not Found Paper", source_status="not_found"),
            Reference(id=3, title="API Error Paper", source_status="api_error"),
            Reference(id=4, title="Pending Paper", source_status="pending"),
        ]
        state = PipelineState(
            session_id="test",
            manuscript=ParsedManuscript(filename="test.docx"),
            references=refs,
        )
        report_path = tmp_path / "report.docx"
        generate_report(state, report_path)

        report_doc = docx.Document(str(report_path))
        full_text = "\n".join(p.text for p in report_doc.paragraphs)
        assert "Found in databases: 1" in full_text
        assert "Not found: 1" in full_text
        assert "API errors: 1" in full_text
