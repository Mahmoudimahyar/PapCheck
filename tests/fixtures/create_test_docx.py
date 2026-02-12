"""Helper to programmatically create test DOCX fixtures."""

from pathlib import Path

import docx


def create_simple_10refs(output_path: Path) -> Path:
    """Create a DOCX with 10 numbered biomedical references."""
    doc = docx.Document()

    doc.add_heading("Effect of Novel Biomarkers on Cancer Prognosis", level=1)

    doc.add_heading("Abstract", level=2)
    doc.add_paragraph(
        "This study evaluates novel biomarkers for cancer prognosis. "
        "We reviewed recent literature [1-3] and conducted meta-analysis."
    )

    doc.add_heading("Introduction", level=2)
    doc.add_paragraph(
        "Cancer biomarkers have transformed clinical practice [1]. "
        "Recent advances in genomics [2, 3] and proteomics [4] have "
        "enabled precision medicine approaches. Several studies [5, 7] "
        "have shown improved outcomes with biomarker-guided therapy. "
        "The BRCA pathway [6] remains a key target. Multi-marker panels "
        "[8-10] show promise for early detection."
    )

    doc.add_heading("Methods", level=2)
    doc.add_paragraph(
        "We performed a systematic review following PRISMA guidelines [1]. "
        "Statistical analysis used methods described by Smith et al. [2]."
    )

    doc.add_heading("Results", level=2)
    doc.add_paragraph(
        "Drug X reduced mortality by 30% [5]. The combination therapy "
        "showed synergistic effects [7, 8]. Biomarker panel accuracy "
        "was 95% [9, 10]."
    )

    doc.add_heading("References", level=2)
    refs = [
        "1. Smith J, Doe A, Lee B. Drug X reduces mortality in elderly "
        "patients: a randomized trial. Lancet. 2020;395:1-10. "
        "doi:10.1016/S0140-6736(20)30001-1 PMID:32456789",

        "2. Jones K, Wang L. CRISPR-based diagnostics for cancer "
        "biomarker detection. Nature Methods. 2019;16(11):1123-1130. "
        "doi:10.1038/s41592-019-0001-1 PMID:31636429",

        "3. Chen X, Park S, Kim H. Proteomic profiling of breast cancer "
        "subtypes. J Proteome Res. 2021;20(3):1456-1468. "
        "doi:10.1021/acs.jproteome.0c00891",

        "4. Brown M, Taylor R. Mass spectrometry in clinical proteomics. "
        "Clin Chem. 2018;64(4):634-642.",

        "5. Williams A, Davis C, Miller E. Long-term outcomes of "
        "biomarker-guided therapy in breast cancer. J Clin Oncol. "
        "2022;40(15):1678-1689. PMID:35731001",

        "6. Garcia F, Martinez D. BRCA1/2 mutations and therapeutic "
        "implications. Cancer Res. 2019;79(12):3021-3030. "
        "doi:10.1158/0008-5472.CAN-19-0001",

        "7. Thompson P, White S. Combination immunotherapy in advanced "
        "melanoma: a phase III trial. N Engl J Med. 2023;388:112-124.",

        "8. Anderson R, Clark J, Lewis M. Multi-cancer early detection "
        "blood test validation study. Science. 2022;376:1-8. "
        "doi:10.1126/science.abc1234",

        "9. Wilson T, Harris N. Circulating tumor DNA as a prognostic "
        "biomarker in colorectal cancer. Gut. 2021;70(5):891-899.",

        "10. Lee S, Kim Y, Park J. Next-generation sequencing panel for "
        "hereditary cancer risk assessment. Genet Med. 2020;22(10):"
        "1592-1600. doi:10.1038/s41436-020-0001-1 PMID:32601389",
    ]
    for ref in refs:
        doc.add_paragraph(ref)

    doc.save(str(output_path))
    return output_path


def create_no_refs_docx(output_path: Path) -> Path:
    """Create a DOCX with no reference section."""
    doc = docx.Document()
    doc.add_heading("A Paper Without References", level=1)
    doc.add_heading("Introduction", level=2)
    doc.add_paragraph("This paper has no reference section at all.")
    doc.add_heading("Methods", level=2)
    doc.add_paragraph("We used standard methods.")
    doc.add_heading("Conclusion", level=2)
    doc.add_paragraph("Results were inconclusive.")
    doc.save(str(output_path))
    return output_path


def create_messy_docx(output_path: Path) -> Path:
    """Create a DOCX with messy formatting."""
    doc = docx.Document()
    doc.add_heading("Messy Manuscript", level=1)
    doc.add_paragraph("Introduction text with some citations [1] [2].")

    doc.add_heading("References", level=2)
    # Messy formatting: missing fields, inconsistent style
    refs = [
        "1. Smith J. Some paper. 2020.",
        "2. Unknown authors, untitled, no date.",
        "3. Jones K et al. A proper reference. Nature. 2021;1:1-5. "
        "doi:10.1038/test-123",
    ]
    for ref in refs:
        doc.add_paragraph(ref)

    doc.save(str(output_path))
    return output_path


if __name__ == "__main__":
    out_dir = Path(__file__).parent / "manuscripts"
    out_dir.mkdir(exist_ok=True)
    create_simple_10refs(out_dir / "simple_10refs.docx")
    create_no_refs_docx(out_dir / "no_refs.docx")
    create_messy_docx(out_dir / "messy_formatting.docx")
